import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCX_PROJECT = '/Users/avinashsharma/msme_credit_risk_assistant/BTech_Project_Report_MSME_Credit_Risk_Assistant.docx'
DOCX_DESKTOP = '/Users/avinashsharma/Desktop/BTech_Project_Report_MSME_Credit_Risk_Assistant.docx'
LOGO_PATH = '/Users/avinashsharma/msme_credit_risk_assistant/assets/bvdu_logo.png'
DIAG_DIR = '/Users/avinashsharma/msme_credit_risk_assistant/assets/diagrams'

doc = Document()

# Set standard styles
def set_margins(section):
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.0)

# Configure default style (Times New Roman, 12pt)
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
font.color.rgb = RGBColor(0, 0, 0)
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.space_after = Pt(6)

def add_title(text, size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12, space_before=12):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading1(text):
    # Chapter Title: 18pt, Bold, Centered, All Caps
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text.upper())
    run.font.name = 'Times New Roman'
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading2(text):
    # Section Heading: 12pt, Bold, Left Aligned, Capitalized
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text.upper())
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading3(text):
    # Section Subheading: 12pt, Bold, Left Aligned, Title Case
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_body(text, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.25)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    return p

def add_bullet(bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(4)
    run_b = p.add_run(bold_prefix + " ")
    run_b.font.name = 'Times New Roman'
    run_b.font.size = Pt(12)
    run_b.font.bold = True
    run_t = p.add_run(text)
    run_t.font.name = 'Times New Roman'
    run_t.font.size = Pt(12)
    return p

def style_table(table, col_widths, headers, data):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = hdr_cells[i].paragraphs[0].runs[0]
        run.font.name = 'Times New Roman'
        run.font.bold = True
        run.font.size = Pt(10.5)
        # Shading
        shading = parse_xml(r'<w:shd {} w:fill="E2E8F0"/>'.format(nsdecls('w')))
        hdr_cells[i]._tc.get_or_add_tcPr().append(shading)
        
    for r_idx, row_data in enumerate(data):
        row_cells = table.add_row().cells
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            row_cells[c_idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 or len(str(val)) > 15 else WD_ALIGN_PARAGRAPH.CENTER
            run = row_cells[c_idx].paragraphs[0].runs[0]
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            if r_idx % 2 == 1:
                shd = parse_xml(r'<w:shd {} w:fill="F8FAFC"/>'.format(nsdecls('w')))
                row_cells[c_idx]._tc.get_or_add_tcPr().append(shd)

    # Set column widths
    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = Inches(w)

print("Helper functions initialized.")
